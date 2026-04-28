"""Helpers for validating and parsing AI assistant document uploads."""

from __future__ import annotations

import csv
import io
import os
from dataclasses import dataclass
from typing import IO, Any


AI_DOCUMENT_ALLOWED_EXTENSIONS = (
    'pdf',
    'txt',
    'csv',
    'jpg',
    'jpeg',
    'png',
    'doc',
    'docx',
    'xls',
    'xlsx',
)

AI_DOCUMENT_ALLOWED_CONTENT_TYPES = frozenset(
    {
        'application/msword',
        'application/octet-stream',
        'application/pdf',
        'application/vnd.ms-excel',
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        'image/jpeg',
        'image/png',
        'text/csv',
        'text/plain',
    }
)

TABULAR_DOCUMENT_EXTENSIONS = frozenset({'csv', 'xls', 'xlsx'})
GENERIC_CONTENT_TYPES = frozenset({'', 'application/octet-stream'})

MAX_TABULAR_ROWS = 500
MAX_TABULAR_TEXT_CHARS = 40000

ROW_TRUNCATION_WARNING = '[WARNING: Data truncated at 500 rows to preserve context window]'
TEXT_TRUNCATION_WARNING = '[WARNING: Markdown output truncated to preserve context window]'


@dataclass(frozen=True)
class TabularParseResult:
    text: str
    rows_included: int
    sheet_count: int
    truncated: bool
    warnings: tuple[str, ...]
    preview: tuple[dict[str, Any], ...]


def get_document_extension(filename: str | None) -> str:
    return os.path.splitext(str(filename or ''))[1].lower().lstrip('.')


def is_tabular_document(*, filename: str | None, content_type: str | None) -> bool:
    extension = get_document_extension(filename)
    if extension in TABULAR_DOCUMENT_EXTENSIONS:
        return True

    normalized_content_type = str(content_type or '').strip().lower()
    return normalized_content_type in {
        'application/vnd.ms-excel',
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        'text/csv',
    }


def validate_ai_document_upload(*, filename: str | None, content_type: str | None) -> None:
    extension = get_document_extension(filename)
    if not extension or extension not in AI_DOCUMENT_ALLOWED_EXTENSIONS:
        allowed = ', '.join(AI_DOCUMENT_ALLOWED_EXTENSIONS)
        raise ValueError(f'Unsupported file extension. Allowed extensions: {allowed}.')

    normalized_content_type = str(content_type or '').strip().lower()
    if normalized_content_type in GENERIC_CONTENT_TYPES:
        return

    if normalized_content_type not in AI_DOCUMENT_ALLOWED_CONTENT_TYPES:
        raise ValueError(
            'Unsupported file content type. Allowed uploads include PDF, image, text, Word, CSV, and Excel files.'
        )


def parse_tabular_document(file_obj: IO[bytes], *, filename: str, content_type: str | None = None) -> TabularParseResult:
    extension = get_document_extension(filename)
    if extension not in TABULAR_DOCUMENT_EXTENSIONS:
        raise ValueError(f'Tabular parser does not support .{extension or "unknown"} files.')

    file_obj.seek(0)

    if extension == 'csv':
        sheets, total_rows, truncated = _parse_csv_sheets(file_obj)
    elif extension == 'xlsx':
        sheets, total_rows, truncated = _parse_xlsx_sheets(file_obj)
    else:
        sheets, total_rows, truncated = _parse_xls_sheets(file_obj)

    markdown_sections: list[str] = []
    preview_rows: list[dict[str, Any]] = []
    warnings: list[str] = []

    for sheet in sheets:
        if sheet['sheet_name']:
            markdown_sections.append(f"## Sheet: {sheet['sheet_name']}")
        markdown_sections.append(sheet['markdown'])
        preview_rows.append(
            {
                'type': 'table',
                'sheet': sheet['sheet_name'],
                'rows_included': sheet['rows_included'],
                'columns': sheet['columns'],
            }
        )

    if truncated:
        warnings.append(ROW_TRUNCATION_WARNING)

    text = '\n\n'.join(section for section in markdown_sections if section).strip()
    if warnings:
        text = f'{text}\n\n' + '\n'.join(warnings)

    if len(text) > MAX_TABULAR_TEXT_CHARS:
        cut = text[:MAX_TABULAR_TEXT_CHARS]
        cut = cut.rsplit('\n', 1)[0] or cut
        warnings.append(TEXT_TRUNCATION_WARNING)
        text = f'{cut}\n\n{TEXT_TRUNCATION_WARNING}'
        truncated = True

    return TabularParseResult(
        text=text,
        rows_included=total_rows,
        sheet_count=len(sheets),
        truncated=truncated,
        warnings=tuple(dict.fromkeys(warnings)),
        preview=tuple(preview_rows[:10]),
    )


def _parse_csv_sheets(file_obj: IO[bytes]) -> tuple[list[dict[str, Any]], int, bool]:
    raw_bytes = file_obj.read()
    try:
        decoded = raw_bytes.decode('utf-8-sig')
    except UnicodeDecodeError:
        decoded = raw_bytes.decode('latin-1')

    reader = csv.reader(io.StringIO(decoded))
    rows = [list(row) for row in reader]
    sheet, rows_included, truncated = _sheet_from_rows('CSV', rows, MAX_TABULAR_ROWS)
    return ([sheet] if sheet else []), rows_included, truncated


def _parse_xlsx_sheets(file_obj: IO[bytes]) -> tuple[list[dict[str, Any]], int, bool]:
    from openpyxl import load_workbook

    workbook = load_workbook(file_obj, read_only=True, data_only=True)
    sheets: list[dict[str, Any]] = []
    remaining_rows = MAX_TABULAR_ROWS
    total_rows = 0
    truncated = False

    try:
        for worksheet in workbook.worksheets:
            rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            sheet, rows_included, sheet_truncated = _sheet_from_rows(worksheet.title, rows, remaining_rows)
            if sheet:
                sheets.append(sheet)
                total_rows += rows_included
                remaining_rows = max(0, remaining_rows - rows_included)
            if sheet_truncated:
                truncated = True
                break
            if remaining_rows <= 0:
                truncated = True
                break
    finally:
        workbook.close()

    return sheets, total_rows, truncated


def _parse_xls_sheets(file_obj: IO[bytes]) -> tuple[list[dict[str, Any]], int, bool]:
    import xlrd

    workbook = xlrd.open_workbook(file_contents=file_obj.read(), on_demand=True)
    sheets: list[dict[str, Any]] = []
    remaining_rows = MAX_TABULAR_ROWS
    total_rows = 0
    truncated = False

    try:
        for sheet_ref in workbook.sheets():
            rows = [sheet_ref.row_values(row_index) for row_index in range(sheet_ref.nrows)]
            sheet, rows_included, sheet_truncated = _sheet_from_rows(sheet_ref.name, rows, remaining_rows)
            if sheet:
                sheets.append(sheet)
                total_rows += rows_included
                remaining_rows = max(0, remaining_rows - rows_included)
            if sheet_truncated:
                truncated = True
                break
            if remaining_rows <= 0:
                truncated = True
                break
    finally:
        workbook.release_resources()

    return sheets, total_rows, truncated


def _sheet_from_rows(
    sheet_name: str,
    rows: list[list[Any]],
    row_budget: int,
) -> tuple[dict[str, Any] | None, int, bool]:
    if not rows:
        return None, 0, False

    normalized_rows = [_normalize_row(row) for row in rows]
    header = normalized_rows[0]
    data_rows = normalized_rows[1:]

    if row_budget < len(data_rows):
        data_rows = data_rows[:row_budget]
        truncated = True
    else:
        truncated = False

    if not any(header):
        column_count = max((len(row) for row in data_rows), default=1)
        header = [f'Column {index + 1}' for index in range(column_count)]

    column_count = max(len(header), *(len(row) for row in data_rows), 0)
    header = _pad_row(header, column_count)
    padded_rows = [_pad_row(row, column_count) for row in data_rows]

    markdown = _rows_to_markdown(header, padded_rows)
    return (
        {
            'sheet_name': sheet_name,
            'rows_included': len(padded_rows),
            'columns': header,
            'markdown': markdown,
        },
        len(padded_rows),
        truncated,
    )


def _normalize_row(row: list[Any]) -> list[str]:
    return [_normalize_cell(value) for value in row]


def _normalize_cell(value: Any) -> str:
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).replace('\r', ' ').replace('\n', ' ').strip()
    text = text.replace('|', '\\|')
    return text


def _pad_row(row: list[str], column_count: int) -> list[str]:
    if len(row) >= column_count:
        return row[:column_count]
    return row + [''] * (column_count - len(row))


def _rows_to_markdown(header: list[str], rows: list[list[str]]) -> str:
    divider = ['---'] * len(header)
    lines = [
        f"| {' | '.join(header)} |",
        f"| {' | '.join(divider)} |",
    ]
    for row in rows:
        lines.append(f"| {' | '.join(row)} |")
    return '\n'.join(lines)
