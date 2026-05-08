"""Source-row loading helpers for Golden Schema ETL dry runs."""

from __future__ import annotations

import csv
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .contracts import SourceBinding


@dataclass(frozen=True)
class LoadedSourceRow:
    """One normalized source row loaded from disk for dry-run processing."""

    entity: str
    source_path: str
    source_sheet: str | None
    source_row_number: int
    payload: dict[str, Any]


def load_source_rows(source: SourceBinding, *, manifest_dir: Path) -> list[LoadedSourceRow]:
    """Load source rows for a single manifest binding."""

    resolved_path = Path(source.path)
    if not resolved_path.is_absolute():
        resolved_path = (manifest_dir / source.path).resolve()

    if source.format == "json":
        return _load_json_rows(source, resolved_path)
    if source.format == "csv":
        return _load_csv_rows(source, resolved_path)
    if source.format == "xlsx":
        return _load_xlsx_rows(source, resolved_path)
    if source.format == "database_export":
        return _load_json_rows(source, resolved_path)

    raise ValueError(f"Unsupported source format: {source.format}")


def _load_json_rows(source: SourceBinding, resolved_path: Path) -> list[LoadedSourceRow]:
    payload = json.loads(resolved_path.read_text(encoding="utf-8"))
    rows: list[LoadedSourceRow] = []

    if isinstance(payload, list):
        iterable = [(source.sheet, payload)]
    elif isinstance(payload, dict) and isinstance(payload.get("rows"), list):
        iterable = [(payload.get("sheet") or source.sheet, payload["rows"])]
    elif isinstance(payload, dict):
        iterable = [(sheet_name, rows_payload) for sheet_name, rows_payload in payload.items()]
    else:
        raise ValueError(f"Unsupported JSON ETL payload shape for {resolved_path}.")

    for sheet_name, row_payloads in iterable:
        if not isinstance(row_payloads, list):
            raise ValueError(f"JSON source sheet {sheet_name!r} must contain a row list.")
        for row_number, row in enumerate(row_payloads, start=1):
            if not isinstance(row, dict):
                raise ValueError(f"JSON source row #{row_number} in {resolved_path} must be an object.")
            rows.append(
                LoadedSourceRow(
                    entity=source.entity,
                    source_path=source.path,
                    source_sheet=str(sheet_name) if sheet_name else None,
                    source_row_number=row_number,
                    payload={str(key): value for key, value in row.items()},
                )
            )

    return rows


def _load_csv_rows(source: SourceBinding, resolved_path: Path) -> list[LoadedSourceRow]:
    with resolved_path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        rows = []
        for row_number, row in enumerate(reader, start=2):
            rows.append(
                LoadedSourceRow(
                    entity=source.entity,
                    source_path=source.path,
                    source_sheet=source.sheet or "CSV",
                    source_row_number=row_number,
                    payload={str(key): value for key, value in row.items() if key is not None},
                )
            )
    return rows


def _load_xlsx_rows(source: SourceBinding, resolved_path: Path) -> list[LoadedSourceRow]:
    workbook = load_workbook(resolved_path, read_only=True, data_only=True)
    rows: list[LoadedSourceRow] = []
    try:
        worksheets = [workbook[source.sheet]] if source.sheet else list(workbook.worksheets)
        for worksheet in worksheets:
            header: list[str] | None = None
            for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                normalized_row = [_normalize_cell(value) for value in row]
                if row_number == 1:
                    header = [value or f"column_{index + 1}" for index, value in enumerate(normalized_row)]
                    continue
                if header is None:
                    raise ValueError(f"Worksheet {worksheet.title} in {resolved_path} is missing a header row.")
                payload = {
                    header[index]: normalized_row[index] if index < len(normalized_row) else ""
                    for index in range(len(header))
                }
                rows.append(
                    LoadedSourceRow(
                        entity=source.entity,
                        source_path=source.path,
                        source_sheet=worksheet.title,
                        source_row_number=row_number,
                        payload=payload,
                    )
                )
    finally:
        workbook.close()
    return rows


def _normalize_cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return value
